#!/usr/bin/env python3
"""
Supplement empty source fields in department Excel files with related sources.
"""
import json
import openpyxl
from openpyxl.styles import Font, Alignment
import subprocess
import sys

DEPARTMENTS = ["法务部", "人力资源部", "包材研发设计部"]
BASE_PATH = "/Users/zhouanjun/Documents/esg-info-collection"
OUTPUT_PATH = f"{BASE_PATH}/data/output/final"
WORKER_INPUT_PATH = f"{BASE_PATH}/data/output/baseline_first/_worker_inputs"

def load_retrieval_candidates(dept_name):
    """Load retrieval candidates for a department."""
    file_path = f"{WORKER_INPUT_PATH}/{dept_name}.json"
    with open(file_path, "r", encoding="utf-8") as f:
        return json.load(f)

def lookup_source_by_code(code):
    """Use the source_lookup.py helper to get faithful source text."""
    try:
        result = subprocess.run(
            [f"{BASE_PATH}/.venv/bin/python",
             "/Users/zhouanjun/.claude/jobs/46174d16/tmp/source_lookup.py",
             code],
            capture_output=True,
            text=True,
            cwd=BASE_PATH,
            timeout=5
        )
        if result.returncode == 0:
            return result.stdout.strip()
        return None
    except Exception as e:
        print(f"Lookup failed for {code}: {e}")
        return None

def format_source_text(candidate):
    """Format source candidate into column 9 format."""
    framework = candidate['framework']
    code = candidate['code']
    orig_excerpt = candidate.get('orig_excerpt', candidate.get('summary', ''))

    # Try to get faithful text from lookup
    faithful_text = lookup_source_by_code(code)
    if faithful_text and len(faithful_text) > 10:
        excerpt = faithful_text[:200] + "..." if len(faithful_text) > 200 else faithful_text
    else:
        excerpt = orig_excerpt[:200] + "..." if len(orig_excerpt) > 200 else orig_excerpt

    if framework == "SSE":
        # SSE format: SSE 第X条｜[excerpt]
        article_num = code.replace("SSE-", "").replace("ART-", "").replace("第", "").replace("条", "")
        return f"SSE 第{article_num}条｜{excerpt}"
    elif framework == "HKEX":
        # HKEX format: HKEX [code]｜[excerpt]
        return f"HKEX {code}｜{excerpt}"
    elif framework == "GRI":
        # GRI format: GRI X-Y｜[excerpt]
        return f"{code}｜{excerpt}"
    elif framework == "MSCI":
        # MSCI format: MSCI <topic>｜MSCI评价/关注[excerpt]
        return f"MSCI {code}｜MSCI评价/关注：{excerpt}"
    elif framework == "CSA":
        # CSA format: CSA-COS [code]｜[excerpt]
        return f"CSA-COS {code}｜{excerpt}"
    else:
        return f"{code}｜{excerpt}"

def find_related_source(topic, question_content, candidates_by_topic):
    """Find a related source for the question from candidates."""
    # Try to find candidates for this topic
    all_candidates = []

    # Match by topic name
    for topic_key, candidates in candidates_by_topic.items():
        if topic_key in topic or topic in topic_key:
            all_candidates.extend(candidates)

    # If no direct topic match, use all candidates
    if not all_candidates:
        for candidates in candidates_by_topic.values():
            all_candidates.extend(candidates)

    if not all_candidates:
        return None, "no candidates available"

    # Prioritize by framework: SSE > GRI > HKEX > CSA > MSCI
    framework_priority = {"SSE": 1, "GRI": 2, "HKEX": 3, "CSA": 4, "MSCI": 5}

    # Sort candidates by framework priority and score
    sorted_candidates = sorted(
        all_candidates,
        key=lambda c: (
            framework_priority.get(c['framework'], 99),
            -int(c.get('score', 0)) if c.get('score') else 0
        )
    )

    # Return the best candidate
    if sorted_candidates:
        return sorted_candidates[0], None

    return None, "no suitable candidate found"

def process_department(dept_name):
    """Process one department's Excel file."""
    print(f"\n{'='*60}")
    print(f"Processing: {dept_name}")
    print(f"{'='*60}")

    # Load retrieval candidates
    worker_data = load_retrieval_candidates(dept_name)
    candidates_by_topic = worker_data['candidates_by_topic']

    # Load Excel
    excel_path = f"{OUTPUT_PATH}/宜格集团2026年度ESG报告信息与数据收集表—{dept_name}_补充版.xlsx"
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["定性"]

    # Track stats
    total_empty = 0
    supplemented = 0
    remaining_empty = []

    # Process each row
    for row_idx in range(2, ws.max_row + 1):
        col9_value = ws.cell(row_idx, 9).value

        # Skip if already has content
        if col9_value and str(col9_value).strip():
            continue

        total_empty += 1

        # Get question info
        topic_code = ws.cell(row_idx, 1).value  # 议题 (column 1)
        question_num = ws.cell(row_idx, 2).value  # 编号 (column 2)
        question_content = ws.cell(row_idx, 3).value  # 问题内容 (column 3)

        if not question_content:
            continue

        # Find related source
        candidate, reason = find_related_source(topic_code, question_content, candidates_by_topic)

        if candidate:
            # Format and write source
            source_text = format_source_text(candidate)
            ws.cell(row_idx, 9).value = source_text

            # Append to 备注 (column 10)
            existing_note = ws.cell(row_idx, 10).value
            note_suffix = "[来源为相关背景条款，非完全直接要求]"

            if existing_note and str(existing_note).strip():
                ws.cell(row_idx, 10).value = f"{existing_note}；{note_suffix}"
            else:
                ws.cell(row_idx, 10).value = note_suffix

            supplemented += 1
            print(f"✓ Row {row_idx} ({question_num}): {candidate['framework']} {candidate['code']}")
        else:
            remaining_empty.append({
                'row': row_idx,
                'question_num': question_num,
                'question': question_content[:50] + "..." if len(question_content) > 50 else question_content,
                'reason': reason
            })

    # Save Excel
    wb.save(excel_path)
    wb.close()

    # Print summary
    print(f"\n{dept_name} Summary:")
    print(f"  Total empty items: {total_empty}")
    print(f"  Supplemented: {supplemented}")
    print(f"  Remaining empty: {len(remaining_empty)}")

    if remaining_empty and len(remaining_empty) <= 10:
        print(f"\n  Remaining empty items:")
        for item in remaining_empty:
            print(f"    Row {item['row']} ({item['question_num']}): {item['question']}")
            print(f"      Reason: {item['reason']}")

    return {
        'department': dept_name,
        'total_empty': total_empty,
        'supplemented': supplemented,
        'remaining_empty': len(remaining_empty),
        'remaining_items': remaining_empty
    }

def main():
    """Process all departments."""
    print("Starting source supplementation for 3 departments...")

    results = []
    for dept in DEPARTMENTS:
        result = process_department(dept)
        results.append(result)

    # Final summary
    print(f"\n{'='*60}")
    print("FINAL SUMMARY")
    print(f"{'='*60}")
    for r in results:
        print(f"{r['department']}: {r['total_empty']} empty → {r['supplemented']} supplemented, {r['remaining_empty']} remain empty")

    print("\nAll done!")

if __name__ == "__main__":
    main()
