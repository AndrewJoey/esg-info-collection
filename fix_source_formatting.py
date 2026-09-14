#!/usr/bin/env python3
"""
Fix formatting issues in source text (remove double 第第 and 条条).
"""
import openpyxl
import re

DEPARTMENTS = ["法务部", "人力资源部", "包材研发设计部"]
BASE_PATH = "/Users/zhouanjun/Documents/esg-info-collection"
OUTPUT_PATH = f"{BASE_PATH}/data/output/final"

def fix_source_formatting(source_text):
    """Fix formatting issues in source text."""
    if not source_text:
        return source_text

    # Fix SSE 第第X条条 → SSE 第X条
    source_text = re.sub(r'SSE 第第(\d+)条条', r'SSE 第\1条', source_text)

    # Fix any other double characters
    source_text = re.sub(r'第第', '第', source_text)
    source_text = re.sub(r'条条', '条', source_text)

    return source_text

def process_department(dept_name):
    """Process one department's Excel file."""
    print(f"\nProcessing: {dept_name}")

    # Load Excel
    excel_path = f"{OUTPUT_PATH}/宜格集团2026年度ESG报告信息与数据收集表—{dept_name}_补充版.xlsx"
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["定性"]

    fixed_count = 0

    # Process each row
    for row_idx in range(2, ws.max_row + 1):
        source_value = ws.cell(row_idx, 9).value

        if source_value and str(source_value).strip():
            original = str(source_value)
            fixed = fix_source_formatting(original)

            if fixed != original:
                ws.cell(row_idx, 9).value = fixed
                fixed_count += 1
                if fixed_count <= 3:  # Show first 3 examples
                    print(f"  Row {row_idx}:")
                    print(f"    Before: {original[:80]}...")
                    print(f"    After:  {fixed[:80]}...")

    # Save Excel
    wb.save(excel_path)
    wb.close()

    print(f"  Fixed {fixed_count} rows")
    return fixed_count

def main():
    """Process all departments."""
    print("Fixing source formatting issues...")

    total_fixed = 0
    for dept in DEPARTMENTS:
        fixed = process_department(dept)
        total_fixed += fixed

    print(f"\nTotal fixed: {total_fixed}")

if __name__ == "__main__":
    main()
