"""Tests for the P1.5 Topic x Department normalization workflow.

All fixtures are SYNTHETIC. No real client department names, topic
lists, or workbook contents appear here.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from backend.models import Topic
from scripts.p15_normalize_topic_department import (
    HEADER_DEPARTMENT,
    HEADER_TOPIC,
    normalize_workbook,
    write_outputs,
)

# Synthetic, obviously-fake names.
DEPT_A = "Alpha Dept"
DEPT_B = "Bravo Dept"
TOPIC_X = "Synthetic Topic X"
TOPIC_Y = "Synthetic Topic Y"
TOPIC_Z = "Synthetic Topic Z"


def _make_workbook(path: Path, rows: list[tuple[str | None, str | None]],
                   merged_header: bool = True) -> None:
    """Build a synthetic 2-column workbook.

    ``rows`` are data rows (department, topic); None means a blank cell
    (simulating vertical merge / forward-fill).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(row=1, column=1, value=HEADER_DEPARTMENT)
    ws.cell(row=1, column=2, value=HEADER_TOPIC)
    if merged_header:
        ws.merge_cells("A1:A2")
        ws.merge_cells("B1:B2")
        start = 3
    else:
        start = 2
    for i, (dept, topic) in enumerate(rows):
        r = start + i
        if dept is not None:
            ws.cell(row=r, column=1, value=dept)
        if topic is not None:
            ws.cell(row=r, column=2, value=topic)
    wb.save(path)


@pytest.fixture
def wb_path(tmp_path: Path) -> Path:
    return tmp_path / "synthetic.xlsx"


class TestForwardFill:
    def test_forward_filled_department_groups(self, wb_path):
        _make_workbook(
            wb_path,
            [
                (DEPT_A, TOPIC_X),
                (None, TOPIC_Y),   # blank dept -> forward-fill DEPT_A
                (None, TOPIC_Z),
                (DEPT_B, TOPIC_X),
                (None, TOPIC_Y),   # forward-fill DEPT_B
            ],
        )
        result = normalize_workbook(wb_path)
        by_dept = {}
        for dept, _tid, topic, _row in result.scope_pairs:
            by_dept.setdefault(dept, []).append(topic)
        assert by_dept[DEPT_A] == [TOPIC_X, TOPIC_Y, TOPIC_Z]
        assert by_dept[DEPT_B] == [TOPIC_X, TOPIC_Y]

    def test_multiple_topics_under_one_department(self, wb_path):
        _make_workbook(
            wb_path,
            [(DEPT_A, TOPIC_X), (None, TOPIC_Y), (None, TOPIC_Z)],
        )
        result = normalize_workbook(wb_path)
        assert len(result.departments) == 1
        assert len(result.scope_pairs) == 3

    def test_one_topic_under_multiple_departments(self, wb_path):
        _make_workbook(
            wb_path,
            [(DEPT_A, TOPIC_X), (DEPT_B, TOPIC_X)],
        )
        result = normalize_workbook(wb_path)
        # Same topic name -> one unique topic, two relationships.
        assert len(result.topics) == 1
        assert len(result.scope_pairs) == 2
        depts = {d for d, _t, _n, _r in result.scope_pairs}
        assert depts == {DEPT_A, DEPT_B}


class TestDetection:
    def test_duplicate_pair_detection(self, wb_path):
        _make_workbook(
            wb_path,
            [(DEPT_A, TOPIC_X), (None, TOPIC_X)],  # duplicate A x X
        )
        result = normalize_workbook(wb_path)
        assert len(result.duplicate_pairs) == 1
        assert result.duplicate_pairs[0] == (DEPT_A, TOPIC_X)
        # Duplicate excluded from the scope mapping.
        assert len(result.scope_pairs) == 1

    def test_blank_rows(self, wb_path):
        _make_workbook(
            wb_path,
            [(DEPT_A, TOPIC_X), (None, None), (None, TOPIC_Y)],
        )
        result = normalize_workbook(wb_path)
        assert result.blank_rows == 1
        # Blank row does not break forward-fill.
        assert len(result.scope_pairs) == 2

    def test_topic_with_missing_department(self, wb_path):
        # First data row has a blank department and nothing to fill from.
        _make_workbook(
            wb_path,
            [(None, TOPIC_X), (DEPT_A, TOPIC_Y)],
        )
        result = normalize_workbook(wb_path)
        assert result.invalid_rows == 1
        statuses = {a.validation_status for a in result.audit_rows}
        assert "invalid_missing_department" in statuses

    def test_department_with_missing_topic(self, wb_path):
        _make_workbook(
            wb_path,
            [(DEPT_A, TOPIC_X), (None, None)],
        )
        # blank/blank counts as blank_row, not invalid; test dept + blank topic:
        _make_workbook(
            wb_path,
            [(DEPT_A, None), (DEPT_A, TOPIC_X)],
        )
        result = normalize_workbook(wb_path)
        assert result.invalid_rows == 1
        statuses = [a.validation_status for a in result.audit_rows]
        assert "invalid_missing_topic" in statuses


class TestWhitespace:
    def test_whitespace_normalization(self, wb_path):
        _make_workbook(
            wb_path,
            [("  " + DEPT_A + " ", "\t" + TOPIC_X + "\n")],
        )
        result = normalize_workbook(wb_path)
        dept, _tid, topic, _row = result.scope_pairs[0]
        assert dept == DEPT_A
        assert topic == TOPIC_X


class TestTopicModel:
    def test_deterministic_topic_ids(self, wb_path):
        _make_workbook(wb_path, [(DEPT_A, TOPIC_X)])
        r1 = normalize_workbook(wb_path)
        r2 = normalize_workbook(wb_path)
        assert list(r1.topics.keys()) == list(r2.topics.keys())
        assert all(tid.startswith("TOPIC-") for tid in r1.topics)

    def test_topics_validate_against_p1_model(self, wb_path):
        _make_workbook(
            wb_path,
            [(DEPT_A, TOPIC_X), (None, TOPIC_Y), (DEPT_B, TOPIC_Z)],
        )
        result = normalize_workbook(wb_path)
        assert result.topic_model_errors == []
        # Every topic id/name is a valid canonical Topic.
        for topic_id, name in result.topics.items():
            t = Topic(topic_id=topic_id, name=name, level=1)
            assert t.level == 1
            assert t.parent_topic_id is None


class TestOutputs:
    def test_output_row_counts_and_provenance(self, wb_path, tmp_path):
        _make_workbook(
            wb_path,
            [
                (DEPT_A, TOPIC_X),
                (None, TOPIC_Y),
                (DEPT_B, TOPIC_X),  # duplicate topic name, new dept
            ],
        )
        result = normalize_workbook(wb_path)
        out_dir = tmp_path / "out"
        written = write_outputs(result, out_dir)

        names = {p.name for p in written}
        assert names == {
            "topic_master.csv",
            "topic_department_scope_mapping.csv",
            "source_row_audit.csv",
            "validation_summary.json",
            "validation_report.md",
        }

        import csv as _csv

        with (out_dir / "topic_master.csv").open(encoding="utf-8") as f:
            master = list(_csv.DictReader(f))
        assert len(master) == 2  # TOPIC_X, TOPIC_Y unique
        assert {m["topic_name"] for m in master} == {TOPIC_X, TOPIC_Y}
        assert all(m["level"] == "1" for m in master)
        assert all(m["parent_topic_id"] == "" for m in master)

        with (out_dir / "topic_department_scope_mapping.csv").open(
            encoding="utf-8"
        ) as f:
            scope = list(_csv.DictReader(f))
        assert len(scope) == 3  # A-X, A-Y, B-X
        # Provenance: source_row must be populated and numeric.
        assert all(row["source_row"].isdigit() for row in scope)
        assert all(row["source_file"] for row in scope)

    def test_source_row_audit_traces_every_row(self, wb_path, tmp_path):
        _make_workbook(
            wb_path,
            [(DEPT_A, TOPIC_X), (None, None), (None, TOPIC_Y)],
        )
        result = normalize_workbook(wb_path)
        out_dir = tmp_path / "out"
        write_outputs(result, out_dir)

        import csv as _csv

        with (out_dir / "source_row_audit.csv").open(encoding="utf-8") as f:
            audit = list(_csv.DictReader(f))
        # One audit row per processed source row (2 data + 1 blank).
        assert len(audit) == 3
        assert all(a["source_row"].isdigit() for a in audit)
        assert all(a["source_sheet"] == "Sheet1" for a in audit)
