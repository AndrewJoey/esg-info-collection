#!/usr/bin/env python3
"""Generate high-quality ESG collection sheet for 包材研发设计部."""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

# Define source mappings based on the baseline data review
SOURCE_MAPPINGS = {
    "D001": "【GRI 301-1】GRI 301-1 材料使用量：报告报告期内用于生产和包装组织主要产品和服务的材料总重量或总体积，分为：i. 不可再生材料用量；ii. 可再生材料用量。建议：计入总材料用量的类型应包括用于包装用途的材料，如纸、纸板和塑料。",
    "D002": "【GRI 301-2】GRI 301-2 再生投入材料使用量：报告用于制造组织主要产品和服务的再生投入材料所占百分比。计算公式：再生投入材料使用总量 ÷ 投入材料使用总量 × 100。",
    "D003": "",  # No direct source found
    "D004": "",  # No direct source found
    "D005": "【SSE第三十七条】上海证券交易所上市公司可持续发展报告指引第三十七条：披露主体应当披露报告期内循环经济的具体情况，包括报告期内为实现循环经济而采取的具体措施，包括节省资源、提高资源利用率、使用可再生资源、预防和减少废弃物的产生以及回收利用废弃物等；报告期内公司在实现循环经济目标方面取得的具体进展及成效，包括废弃物的回收及综合利用情况（含废弃物循环利用量）。",
    "D006": "",  # Derived metric
    "D007": "",  # No direct source found
    "D008": "",  # No direct source found
    "AI-D-01": "【GRI 301-3】GRI 301-3 回收再利用的产品及其包装材料：报告各产品类别中，回收再利用的产品及其包装材料所占百分比；以及本项披露数据的收集方式。计算公式：报告期内回收再利用的产品及其包装材料 ÷ 报告期内售出的产品 × 100（剔除次品与召回品）。",
}

QUAL_SOURCE_MAPPINGS = {
    "Q003": "【MSCI包装材料与废弃物】MSCI在包装材料与废弃物(Packaging Material & Waste)关键议题下评价公司减少包装环境影响策略的全面性，例如通过减少废弃物、包装再设计、以及与供应商或客户的协作。",
    "Q005": "【MSCI包装材料与废弃物】MSCI在包装材料与废弃物关键议题下关注的风险包括：因包装材料与废弃物相关法规改革而失去市场准入；因消费者偏好改变而导致收入损失；因重新设计包装及遵守生产者责任法规而增加成本。",
    "Q007": "【MSCI包装材料与废弃物】MSCI在包装材料与废弃物关键议题下评价公司在包装成分方面的目标(如提高再生成分、可再生成分，或减少难回收塑料)，以及在包装材料循环性或减废方面的目标。",
    "Q011": "【SSE第三十四条】上海证券交易所上市公司可持续发展报告指引第三十四条：披露主体应当集约、高效利用能源、水、原材料等资源，加强资源使用过程节约管理，推动生产、流通过程的减量化、再利用、再循环。",
    "Q014": "【SSE第三十七条】上海证券交易所上市公司可持续发展报告指引第三十七条：披露主体应当披露报告期内循环经济的具体情况，包括报告期内为实现循环经济而采取的具体措施，包括节省资源、提高资源利用率、使用可再生资源、预防和减少废弃物的产生以及回收利用废弃物等。",
    "Q015": "【SSE第三十七条】上海证券交易所上市公司可持续发展报告指引第三十七条：披露主体应当披露报告期内循环经济的具体情况，包括为实现循环经济而制定的具体目标和计划；报告期内公司在实现循环经济目标方面取得的具体进展及成效。",
    "Q031": "【SSE第五十四条】上海证券交易所上市公司可持续发展报告指引第五十四条：披露主体在经营活动中，应当遵循自愿、公平、等价有偿、诚实信用的原则，遵守社会公德、商业道德，不得通过贿赂等非法活动谋取不正当利益，不得侵犯他人的商标权、专利权和著作权等知识产权，不得从事不正当竞争行为。",
    "AI-Q-01": "【MSCI包装材料与废弃物】MSCI在包装材料与废弃物关键议题下评价公司面向消费者的减废教育项目范围，例如回收点(take-back)或回收指引(recycling instructions)。",
}

def load_baseline_data():
    """Load baseline qualitative and quantitative data."""
    qual_df = pd.read_csv("/Users/zhouanjun/Documents/esg-info-collection/data/output/baseline_first/包材研发设计部/qualitative_draft.csv")
    quant_df = pd.read_csv("/Users/zhouanjun/Documents/esg-info-collection/data/output/baseline_first/包材研发设计部/quantitative_draft.csv")
    return qual_df, quant_df

def assign_dimension(topic, qid):
    """Assign dimension based on question pattern."""
    if qid in ["Q001", "Q002", "Q009", "Q010", "Q017", "Q018", "Q026", "Q027"]:
        return "治理"
    elif qid in ["Q003", "Q011", "Q019", "Q028"]:
        return "战略"
    elif qid in ["Q004", "Q005", "Q006", "Q012", "Q013", "Q014", "Q020", "Q021", "Q022", "Q029", "Q030", "Q031"]:
        return "影响、风险和机遇管理"
    elif qid in ["Q007", "Q015", "Q023", "Q032"]:
        return "指标与目标"
    else:
        return "案例实践"

def create_merged_sheet():
    """Create merged sheet with all questions and metrics."""
    qual_df, quant_df = load_baseline_data()

    rows = []
    idx = 1

    # Process qualitative questions
    for _, row in qual_df.iterrows():
        topic = row['议题']
        qid = row['编号']
        question = row['问题内容']
        guidance = row['填写说明口径'] if pd.notna(row['填写说明口径']) else ""
        source = QUAL_SOURCE_MAPPINGS.get(qid, "")
        dimension = assign_dimension(topic, qid)

        rows.append({
            "编号": idx,
            "议题": topic,
            "维度": dimension,
            "问题内容/数据项": question,
            "定性/定量": "定性",
            "填写说明/口径": guidance,
            "来源及原文(中文)": source,
            "备注": ""
        })
        idx += 1

    # Process quantitative metrics
    for _, row in quant_df.iterrows():
        topic = row['议题']
        did = row['编号']
        metric = row['数据项']
        calc = row['填写说明口径'] if pd.notna(row['填写说明口径']) else ""
        source = SOURCE_MAPPINGS.get(did, "")

        rows.append({
            "编号": idx,
            "议题": topic,
            "维度": "指标与目标",
            "问题内容/数据项": metric,
            "定性/定量": "定量",
            "填写说明/口径": calc,
            "来源及原文(中文)": source,
            "备注": ""
        })
        idx += 1

    return pd.DataFrame(rows)

def style_worksheet(ws):
    """Apply styling to worksheet similar to reference file."""
    # Header style
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")

    # Apply header style
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Set column widths
    ws.column_dimensions['A'].width = 8   # 编号
    ws.column_dimensions['B'].width = 18  # 议题
    ws.column_dimensions['C'].width = 22  # 维度
    ws.column_dimensions['D'].width = 60  # 问题内容/数据项
    ws.column_dimensions['E'].width = 10  # 定性/定量
    ws.column_dimensions['F'].width = 40  # 填写说明/口径
    ws.column_dimensions['G'].width = 65  # 来源及原文
    ws.column_dimensions['H'].width = 20  # 备注

    # Data cell style
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = thin_border
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

            # Center align for specific columns
            if cell.column in [1, 5]:  # 编号, 定性/定量
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Freeze header row
    ws.freeze_panes = "A2"

    return ws

def main():
    df = create_merged_sheet()

    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "完整合并_人工+AI补充"

    # Write data
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)

    # Apply styling
    ws = style_worksheet(ws)

    # Save file
    output_path = "/Users/zhouanjun/Documents/esg-info-collection/data/output/final/包材研发设计部_高质量收集表.xlsx"
    wb.save(output_path)

    # Print statistics
    qual_count = len(df[df['定性/定量'] == '定性'])
    quant_count = len(df[df['定性/定量'] == '定量'])
    qual_with_source = len(df[(df['定性/定量'] == '定性') & (df['来源及原文(中文)'].notna()) & (df['来源及原文(中文)'] != '')])
    quant_with_source = len(df[(df['定性/定量'] == '定量') & (df['来源及原文(中文)'].notna()) & (df['来源及原文(中文)'] != '')])

    print(f"\n包材研发设计部收集表已生成")
    print(f"- 定性：{qual_count}项（有来源{qual_with_source}项）")
    print(f"- 定量：{quant_count}项（有来源{quant_with_source}项）")
    print(f"- 总计：{len(df)}项")
    print(f"- 文件：包材研发设计部_高质量收集表.xlsx")

if __name__ == "__main__":
    main()
